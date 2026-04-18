"""PPTX delivery exporter."""

from __future__ import annotations

import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ds_agent.domain.entities.delivery_pack import (
    ArtifactFormat,
    DeliveryArtifact,
    DeliveryPack,
    DeliveryTemplate,
    DesignTheme,
    NarrativeBlocks,
    NarrativeVerification,
    RenderedChart,
)
from ds_agent.infrastructure.exporters.template_registry import TemplateRegistry, coerce_template
from ds_agent.infrastructure.exporters.theme_loader import DEFAULT_THEME, FileSystemThemeLoader


class PptxExporter:
    """Serialize narrative blocks into a presentation deck."""

    supported_format = ArtifactFormat.PPTX

    def __init__(
        self,
        theme: DesignTheme | None = None,
        template_registry: TemplateRegistry | None = None,
        *,
        theme_loader: FileSystemThemeLoader | None = None,
    ) -> None:
        self._theme = theme
        self._template_registry = template_registry
        self._theme_loader = theme_loader

    def export(
        self,
        *,
        artifact: DeliveryArtifact,
        analysis: Mapping[str, Any] | None = None,
        narrative: NarrativeBlocks,
        charts: Sequence[RenderedChart],
        verification: NarrativeVerification | object | None = None,
        output_dir: Path,
        template: DeliveryTemplate | object | None = None,
        pack: object | None = None,
    ) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        analysis_payload = dict(analysis or {})
        if pack is not None:
            analysis_payload.setdefault("confidence", getattr(pack, "confidence", None))
            analysis_payload.setdefault("signed_by", getattr(pack, "signed_by", None))
        template_payload = self._resolve_template(artifact, template)
        verification_payload = _coerce_verification(verification)
        resolved_theme = self._resolve_theme(pack if isinstance(pack, DeliveryPack) else None)

        presentation = self._load_presentation(template_payload)
        for block in narrative.blocks:
            slide = presentation.slides.add_slide(self._content_layout(presentation))
            self._write_title(slide, block.title, theme=resolved_theme)
            self._write_body(slide, block.body_md, theme=resolved_theme)
            self._add_chart_images(slide, block.chart_specs, charts)
        self._append_footer(
            presentation,
            analysis=analysis_payload,
            verification=verification_payload,
            theme=resolved_theme,
        )
        target = output_dir / f"{artifact.artifact_id}.pptx"
        presentation.save(target)
        self._write_preview_manifest(
            target.with_suffix(".preview.json"),
            artifact=artifact,
            narrative=narrative,
            template=template_payload,
        )
        return target

    def _resolve_theme(self, pack: DeliveryPack | None) -> DesignTheme:
        if self._theme is not None:
            return self._theme
        if self._theme_loader is not None:
            return self._theme_loader.resolve(pack=pack)
        return DEFAULT_THEME

    def _resolve_template(
        self,
        artifact: DeliveryArtifact,
        template: DeliveryTemplate | object | None,
    ) -> DeliveryTemplate:
        if isinstance(template, DeliveryTemplate):
            return template
        if template is not None:
            return coerce_template(
                template,
                template_ref=artifact.template_ref,
                format=artifact.format,
                fallback_structure=artifact.content_policy.structure,
                fallback_title=artifact.type.value.replace("_", " ").title(),
            )
        if self._template_registry is not None:
            return coerce_template(
                self._template_registry,
                template_ref=artifact.template_ref,
                format=artifact.format,
                fallback_structure=artifact.content_policy.structure,
                fallback_title=artifact.type.value.replace("_", " ").title(),
            )
        return DeliveryTemplate(
            template_ref=artifact.template_ref,
            family=(
                artifact.template_ref.split("/")[1]
                if "/" in artifact.template_ref
                else artifact.template_ref
            ),
            format=artifact.format,
            title=artifact.type.value.replace("_", " ").title(),
            default_structure=tuple(artifact.content_policy.structure),
        )

    @staticmethod
    def _load_presentation(template: DeliveryTemplate) -> Any:
        if template.template_path:
            candidate = Path(template.template_path)
            if candidate.exists():
                return Presentation(str(candidate))
        return Presentation()

    @staticmethod
    def _content_layout(presentation: Any) -> Any:
        if len(presentation.slide_layouts) > 1:
            return presentation.slide_layouts[1]
        return presentation.slide_layouts[0]

    def _write_title(self, slide: Any, title: str, *, theme: DesignTheme) -> None:
        title_shape = slide.shapes.title
        if title_shape is None:
            title_shape = slide.shapes.add_textbox(
                Inches(0.6),
                Inches(0.4),
                Inches(8.2),
                Inches(0.8),
            )
        frame = title_shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = title
        run.font.name = theme.font_heading
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = self._hex_to_rgb(theme.primary_color)

    def _write_body(self, slide: Any, body_md: str, *, theme: DesignTheme) -> None:
        placeholder = slide.placeholders[1] if len(slide.placeholders) > 1 else None
        if placeholder is None:
            placeholder = slide.shapes.add_textbox(
                Inches(0.8),
                Inches(1.4),
                Inches(8.0),
                Inches(3.5),
            )
        frame = placeholder.text_frame
        frame.clear()
        lines = [line.strip() for line in body_md.splitlines() if line.strip()]
        if not lines:
            lines = ["(no content)"]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.level = 0
            run = paragraph.add_run()
            is_bullet = line.startswith("- ")
            run.text = f"* {line.removeprefix('- ').strip()}" if is_bullet else line
            run.font.name = theme.font_body
            run.font.size = Pt(16)
            run.font.color.rgb = self._hex_to_rgb(theme.secondary_color)

    @staticmethod
    def _add_chart_images(
        slide: Any,
        chart_specs: Sequence[Any],
        charts: Sequence[RenderedChart],
    ) -> None:
        chart_lookup = {chart.chart_id: chart for chart in charts}
        for index, spec in enumerate(chart_specs):
            chart = chart_lookup.get(spec.chart_id)
            if chart is None:
                continue
            left = Inches(5.2)
            top = Inches(1.8 + (index * 1.9))
            slide.shapes.add_picture(str(chart.image_path), left, top, width=Inches(4.0))

    def _append_footer(
        self,
        presentation: Any,
        *,
        analysis: Mapping[str, Any],
        verification: NarrativeVerification,
        theme: DesignTheme,
    ) -> None:
        confidence = analysis.get("confidence")
        signed_by = analysis.get("signed_by", "ds-agent")
        footer_parts = [f"signed_by={signed_by or 'ds-agent'}"]
        if verification.report_id:
            footer_parts.insert(0, f"verifier={verification.report_id}")
        if confidence is not None:
            footer_parts.insert(0, f"confidence={confidence}")
        if verification.flagged_claims:
            footer_parts.append("flags=" + "; ".join(verification.flagged_claims[:2]))
        footer_text = " | ".join(str(part) for part in footer_parts)
        for slide in presentation.slides:
            box = slide.shapes.add_textbox(
                Inches(0.5),
                presentation.slide_height - Inches(0.45),
                Inches(9.0),
                Inches(0.25),
            )
            frame = box.text_frame
            frame.clear()
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = footer_text
            run.font.name = theme.font_body
            run.font.size = Pt(9)
            run.font.color.rgb = self._hex_to_rgb(theme.secondary_color)

    @staticmethod
    def _write_preview_manifest(
        target: Path,
        *,
        artifact: DeliveryArtifact,
        narrative: NarrativeBlocks,
        template: DeliveryTemplate,
    ) -> None:
        slides = []
        for index, block in enumerate(narrative.blocks, start=1):
            body_lines = [line.strip() for line in block.body_md.splitlines() if line.strip()]
            bullets = [
                line.removeprefix("- ").strip()
                for line in body_lines
                if line.startswith("- ")
            ]
            excerpt_source = bullets if bullets else body_lines
            excerpt = " ".join(excerpt_source[:2]).strip()
            slides.append(
                {
                    "index": index,
                    "section": block.section,
                    "title": block.title,
                    "bullets": bullets[:4],
                    "excerpt": excerpt[:240],
                    "chart_count": len(block.chart_specs),
                }
            )

        payload = {
            "artifact_id": artifact.artifact_id,
            "artifact_type": artifact.type.value,
            "audience": artifact.audience.value,
            "template_ref": template.template_ref,
            "slide_count": len(slides),
            "slides": slides,
        }
        target.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    @staticmethod
    def _hex_to_rgb(value: str) -> RGBColor:
        normalized = value.lstrip("#")
        if len(normalized) != 6:
            raise ValueError(f"Expected a 6-digit hex color, got: {value}")
        return cast(RGBColor, RGBColor.from_string(normalized.upper()))


def _coerce_verification(value: NarrativeVerification | object | None) -> NarrativeVerification:
    if isinstance(value, NarrativeVerification):
        return value
    if value is None:
        return NarrativeVerification()
    report_id = getattr(value, "report_id", None)
    flagged_claims = list(getattr(value, "flagged_claims", []) or [])
    if not flagged_claims and hasattr(value, "narrative"):
        flagged_claims = list(value.narrative.flagged_claims or [])
    rejected = not bool(getattr(value, "passed", True))
    message = getattr(value, "message", None)
    return NarrativeVerification(
        report_id=report_id,
        flagged_claims=flagged_claims,
        rejected=rejected,
        notes=[str(message)] if message else [],
    )

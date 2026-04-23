"""Workspace-file exporters (P1-12).

Converts a workspace file into a user-shippable format: PDF, DOCX, HTML,
XLSX, or IPYNB. PDF generation happens in Electron's Chromium via
``webContents.printToPDF``; this module only produces the intermediate HTML
it consumes. All other formats are handled end-to-end here.
"""

from __future__ import annotations

import csv
import html
import json
import shutil
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Any


class ExportFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    HTML = "html"
    XLSX = "xlsx"
    IPYNB = "ipynb"


class ExportError(Exception):
    """Raised when a source file cannot be exported to the requested format."""


@dataclass(frozen=True, slots=True)
class ExportResult:
    output_path: Path
    format: ExportFormat
    # For PDF, the caller (Electron) must render this intermediate HTML.
    intermediate_html_path: Path | None = None


_TABULAR_SOURCES = {".csv", ".tsv"}
_MARKDOWN_SOURCES = {".md", ".markdown"}
_NOTEBOOK_SOURCES = {".ipynb"}
_HTML_SOURCES = {".html", ".htm"}


_SUPPORTED_MATRIX: dict[str, set[ExportFormat]] = {
    ".md": {ExportFormat.PDF, ExportFormat.DOCX, ExportFormat.HTML},
    ".markdown": {ExportFormat.PDF, ExportFormat.DOCX, ExportFormat.HTML},
    ".csv": {ExportFormat.HTML, ExportFormat.XLSX},
    ".tsv": {ExportFormat.HTML, ExportFormat.XLSX},
    ".ipynb": {ExportFormat.PDF, ExportFormat.HTML, ExportFormat.IPYNB},
    ".html": {ExportFormat.PDF, ExportFormat.HTML},
    ".htm": {ExportFormat.PDF, ExportFormat.HTML},
}


def supported_formats(source_suffix: str) -> list[ExportFormat]:
    return sorted(_SUPPORTED_MATRIX.get(source_suffix.lower(), set()))


def export_file(
    source_path: Path,
    output_path: Path,
    export_format: ExportFormat,
) -> ExportResult:
    """Render ``source_path`` as ``export_format`` into ``output_path``.

    The caller is responsible for ensuring both paths are safe (inside the
    workspace / inside a controlled exports dir). We do not re-validate
    here so callers can write to unit-test temp dirs.
    """
    if not source_path.exists():
        raise ExportError(f"Source file not found: {source_path}")
    if not source_path.is_file():
        raise ExportError(f"Source is not a regular file: {source_path}")

    suffix = source_path.suffix.lower()
    allowed = _SUPPORTED_MATRIX.get(suffix)
    if not allowed or export_format not in allowed:
        raise ExportError(
            f"Cannot export {suffix or 'unknown'} -> {export_format.value}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)

    if export_format == ExportFormat.IPYNB:
        if suffix not in _NOTEBOOK_SOURCES:
            raise ExportError("IPYNB export only valid for .ipynb sources")
        shutil.copy2(source_path, output_path)
        return ExportResult(output_path=output_path, format=export_format)

    if export_format == ExportFormat.HTML:
        html_body = _render_html_body(source_path, suffix)
        output_path.write_text(_wrap_html_document(html_body, source_path.stem), encoding="utf-8")
        return ExportResult(output_path=output_path, format=export_format)

    if export_format == ExportFormat.PDF:
        # Render the intermediate HTML; Electron completes the PDF step.
        html_body = _render_html_body(source_path, suffix)
        output_path.write_text(_wrap_html_document(html_body, source_path.stem), encoding="utf-8")
        return ExportResult(
            output_path=output_path,
            format=export_format,
            intermediate_html_path=output_path,
        )

    if export_format == ExportFormat.DOCX:
        if suffix not in _MARKDOWN_SOURCES:
            raise ExportError("DOCX export requires a Markdown source")
        _export_markdown_to_docx(source_path, output_path)
        return ExportResult(output_path=output_path, format=export_format)

    if export_format == ExportFormat.XLSX:
        if suffix not in _TABULAR_SOURCES:
            raise ExportError("XLSX export requires a CSV/TSV source")
        delimiter = "\t" if suffix == ".tsv" else ","
        _export_tabular_to_xlsx(source_path, output_path, delimiter=delimiter)
        return ExportResult(output_path=output_path, format=export_format)

    raise ExportError(f"Unhandled export format: {export_format.value}")


# --- HTML rendering --------------------------------------------------------


def _render_html_body(source_path: Path, suffix: str) -> str:
    if suffix in _MARKDOWN_SOURCES:
        return _markdown_to_html(source_path.read_text(encoding="utf-8"))
    if suffix in _TABULAR_SOURCES:
        delimiter = "\t" if suffix == ".tsv" else ","
        return _csv_to_html_table(source_path, delimiter=delimiter)
    if suffix in _NOTEBOOK_SOURCES:
        return _notebook_to_html(source_path)
    if suffix in _HTML_SOURCES:
        raw = source_path.read_text(encoding="utf-8")
        return raw
    raise ExportError(f"No HTML renderer for {suffix}")


def _markdown_to_html(text: str) -> str:
    try:
        import markdown as md_lib  # type: ignore[import-untyped]
    except ImportError as e:
        raise ExportError(
            "markdown package is required for markdown export"
        ) from e

    rendered: str = md_lib.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists", "toc"],
        output_format="html5",
    )
    return rendered


def _csv_to_html_table(path: Path, *, delimiter: str) -> str:
    rows_html: list[str] = []
    with path.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        for i, row in enumerate(reader):
            cells = "".join(
                f"<{'th' if i == 0 else 'td'}>{html.escape(cell)}</{'th' if i == 0 else 'td'}>"
                for cell in row
            )
            rows_html.append(f"<tr>{cells}</tr>")
    return f"<table>{''.join(rows_html)}</table>"


def _notebook_to_html(path: Path) -> str:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ExportError(f"Invalid notebook JSON: {e}") from e

    cells = data.get("cells") or []
    parts: list[str] = []
    for cell in cells:
        if not isinstance(cell, dict):
            continue
        source = _join_source(cell.get("source"))
        cell_type = cell.get("cell_type")
        if cell_type == "markdown":
            parts.append(_markdown_to_html(source))
        elif cell_type == "code":
            escaped = html.escape(source)
            parts.append(f'<pre class="code-cell"><code>{escaped}</code></pre>')
            for out in cell.get("outputs") or []:
                if not isinstance(out, dict):
                    continue
                text = _join_source(out.get("text"))
                if text:
                    parts.append(
                        f'<pre class="code-output">{html.escape(text)}</pre>'
                    )
    return "\n".join(parts)


def _join_source(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "".join(str(line) for line in value)
    return str(value)


def _wrap_html_document(body: str, title: str) -> str:
    safe_title = html.escape(title)
    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8" />
<title>{safe_title}</title>
<style>
  @page {{ size: A4; margin: 2cm 2.5cm; }}
  body {{
    font-family: "Noto Sans KR", "Malgun Gothic", -apple-system, BlinkMacSystemFont,
                 "Segoe UI", Roboto, "Apple SD Gothic Neo", sans-serif;
    color: #1a1a1a;
    line-height: 1.6;
    max-width: 800px;
    margin: 0 auto;
    padding: 1em;
  }}
  h1, h2, h3 {{ line-height: 1.25; margin-top: 1.5em; }}
  h1 {{ font-size: 1.8em; border-bottom: 2px solid #e5e7eb; padding-bottom: 0.3em; }}
  h2 {{ font-size: 1.4em; }}
  h3 {{ font-size: 1.15em; }}
  code {{
    background: #f4f4f5; padding: 0.15em 0.4em; border-radius: 3px;
    font-family: "SF Mono", Menlo, Consolas, monospace; font-size: 0.9em;
  }}
  pre {{
    background: #f4f4f5; padding: 1em; border-radius: 4px;
    overflow-x: auto; font-size: 0.85em;
  }}
  pre code {{ background: transparent; padding: 0; }}
  pre.code-output {{ background: #fafafa; border-left: 3px solid #a1a1aa; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
  th, td {{ border: 1px solid #d4d4d8; padding: 0.4em 0.7em; text-align: left; }}
  th {{ background: #f4f4f5; font-weight: 600; }}
  img {{ max-width: 100%; height: auto; }}
  blockquote {{
    border-left: 4px solid #d4d4d8; color: #52525b;
    margin: 1em 0; padding-left: 1em;
  }}
</style>
</head>
<body>
{body}
</body>
</html>
"""


# --- DOCX ------------------------------------------------------------------


def _export_markdown_to_docx(source: Path, output: Path) -> None:
    try:
        from docx import Document
    except ImportError as e:
        raise ExportError("python-docx is required for DOCX export") from e

    text = source.read_text(encoding="utf-8")
    document = Document()

    in_code_block = False
    code_buffer: list[str] = []
    in_table = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows = len(table_rows)
        cols = max(len(r) for r in table_rows)
        tbl = document.add_table(rows=rows, cols=cols)
        tbl.style = "Light Grid"
        for r_idx, row in enumerate(table_rows):
            for c_idx in range(cols):
                cell_text = row[c_idx] if c_idx < len(row) else ""
                tbl.rows[r_idx].cells[c_idx].text = cell_text
        in_table = False
        table_rows = []

    def flush_code() -> None:
        nonlocal in_code_block, code_buffer
        if not code_buffer:
            in_code_block = False
            return
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(code_buffer))
        run.font.name = "Consolas"
        code_buffer = []
        in_code_block = False

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            continue
        if in_code_block:
            code_buffer.append(line)
            continue

        if "|" in line and line.strip().startswith("|"):
            # Skip the separator row (|---|---|)
            stripped = line.strip()
            if set(stripped.replace("|", "").strip()) <= {"-", ":", " "}:
                continue
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_rows.append(cells)
            continue
        if in_table:
            flush_table()

        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip() == "":
            document.add_paragraph("")
        else:
            document.add_paragraph(line)

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()

    document.save(str(output))


# --- XLSX ------------------------------------------------------------------


def _export_tabular_to_xlsx(source: Path, output: Path, *, delimiter: str) -> None:
    try:
        from openpyxl import Workbook  # type: ignore[import-untyped]
    except ImportError as e:
        raise ExportError("openpyxl is required for XLSX export") from e

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = source.stem[:31] or "Sheet1"

    with source.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        for row in reader:
            sheet.append([_coerce_cell(value) for value in row])

    workbook.save(output)


def _coerce_cell(value: str) -> str | int | float:
    """Try to coerce numeric strings so Excel treats them as numbers."""
    if value == "":
        return value
    try:
        if "." not in value and "e" not in value.lower():
            return int(value)
        return float(value)
    except ValueError:
        return value

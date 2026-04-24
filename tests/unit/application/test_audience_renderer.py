from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import pytest
from pptx import Presentation

from ds_agent.application.services.audience_renderer import (
    AudienceRenderer,
    NarrativeGenerationRequest,
)
from ds_agent.domain.entities.delivery_pack import (
    ArtifactFormat,
    ArtifactType,
    AudienceKind,
    ContentPolicy,
    DeliveryArtifact,
    DeliveryPack,
    DeliveryTemplate,
    NarrativeVerification,
)
from ds_agent.infrastructure.exporters.pptx_exporter import PptxExporter


class FakeGenerator:
    def __init__(self, payload: dict[str, object]) -> None:
        self.payload = payload
        self.requests: list[NarrativeGenerationRequest] = []

    def generate(self, request: NarrativeGenerationRequest) -> dict[str, object]:
        self.requests.append(request)
        return self.payload


class FakeVerifier:
    def verify(self, artifact, narrative, analysis) -> NarrativeVerification:
        del artifact, analysis
        return NarrativeVerification(
            report_id="vr::test",
            flagged_claims=list(narrative.flagged_claims),
            rejected=False,
        )


class FakeTemplateRegistry:
    def __init__(self, template_path: Path | None = None) -> None:
        self._template_path = template_path

    def load(self, template_ref, *, format, fallback_structure) -> DeliveryTemplate:
        return DeliveryTemplate(
            template_ref=template_ref,
            family="exec_brief",
            format=format,
            title="Executive Brief",
            default_structure=tuple(fallback_structure),
            template_path=str(self._template_path) if self._template_path is not None else None,
        )


class FakeExporter:
    supported_format = ArtifactFormat.MARKDOWN

    def __init__(self) -> None:
        self.output_paths: list[Path] = []

    def export(
        self,
        *,
        artifact,
        narrative,
        output_dir: Path,
        template,
        pack=None,
    ) -> Path:
        del pack, narrative, template
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / f"{artifact.artifact_id}.md"
        path.write_text("rendered", encoding="utf-8")
        self.output_paths.append(path)
        return path


def _pack() -> DeliveryPack:
    artifact = DeliveryArtifact(
        artifact_id="art-exec",
        type=ArtifactType.EXEC_BRIEF,
        audience=AudienceKind.EXECUTIVE,
        format=ArtifactFormat.MARKDOWN,
        content_policy=ContentPolicy(
            structure=["summary", "impact", "decision_needed"],
            tone="decisive",
            technical_detail="minimal",
        ),
        template_ref="tpl/exec_brief/v3",
    )
    return DeliveryPack(
        pack_id="DP-1",
        task_id="TC-2026-001",
        generated_at=datetime(2026, 4, 16, tzinfo=UTC),
        artifacts=[artifact],
    )


def _pptx_pack() -> DeliveryPack:
    artifact = DeliveryArtifact(
        artifact_id="art-exec-pptx",
        type=ArtifactType.EXEC_BRIEF,
        audience=AudienceKind.EXECUTIVE,
        format=ArtifactFormat.PPTX,
        content_policy=ContentPolicy(
            structure=["summary", "impact", "decision_needed"],
            tone="decisive",
            technical_detail="minimal",
        ),
        template_ref="tpl/exec_brief/v3",
    )
    return DeliveryPack(
        pack_id="DP-2",
        task_id="TC-2026-002",
        generated_at=datetime(2026, 4, 16, tzinfo=UTC),
        artifacts=[artifact],
    )


def test_audience_renderer_builds_persona_prompt_and_exports(tmp_path: Path) -> None:
    generator = FakeGenerator(
        {
            "blocks": [
                {
                    "section": "summary",
                    "title": "Summary",
                    "body_md": "Revenue risk increased by 4%.",
                    "citations": ["lineage-1"],
                }
            ],
            "overall_tone": "decisive",
            "flagged_claims": [],
        }
    )
    exporter = FakeExporter()
    renderer = AudienceRenderer(
        generator=generator,
        verifier=FakeVerifier(),
        template_registry=FakeTemplateRegistry(),
        exporters=[exporter],
    )

    rendered = renderer.render(
        pack=_pack(),
        artifact_id="art-exec",
        analysis={"summary": "Revenue risk increased by 4%."},
        output_dir=tmp_path,
    )

    assert rendered.output_path.name == "art-exec.md"
    assert rendered.verifier_report_id == "vr::test"
    assert exporter.output_paths[0].exists()
    assert "C-level decision maker" in generator.requests[0].system_prompt
    assert generator.requests[0].artifact.content_policy.technical_detail == "minimal"


def test_audience_renderer_rejects_invalid_narrative_payload(tmp_path: Path) -> None:
    renderer = AudienceRenderer(
        generator=FakeGenerator({"overall_tone": "decisive"}),
        verifier=FakeVerifier(),
        template_registry=FakeTemplateRegistry(),
        exporters=[FakeExporter()],
    )

    with pytest.raises(ValueError):
        renderer.render(
            pack=_pack(),
            artifact_id="art-exec",
            analysis={"summary": "Revenue risk increased by 4%."},
            output_dir=tmp_path,
        )


def test_audience_renderer_supports_real_pptx_exporter(tmp_path: Path) -> None:
    template_path = tmp_path / "master.pptx"
    Presentation().save(template_path)
    renderer = AudienceRenderer(
        generator=FakeGenerator(
            {
                "blocks": [
                    {
                        "section": "summary",
                        "title": "Summary",
                        "body_md": "Revenue risk increased by 4%.",
                        "citations": ["lineage-1"],
                    }
                ],
                "overall_tone": "decisive",
                "flagged_claims": [],
            }
        ),
        verifier=FakeVerifier(),
        template_registry=FakeTemplateRegistry(template_path),
        exporters=[PptxExporter()],
    )

    rendered = renderer.render(
        pack=_pptx_pack(),
        artifact_id="art-exec-pptx",
        analysis={"summary": "Revenue risk increased by 4%."},
        output_dir=tmp_path,
    )

    assert rendered.output_path.suffix == ".pptx"
    assert rendered.output_path.exists()
    assert len(Presentation(str(rendered.output_path)).slides) == 1

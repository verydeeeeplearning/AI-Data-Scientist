from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from PyPDF2 import PdfReader

from ds_agent.application.dtos.task_contract import (
    BuildDeliveryPackDTO,
    RenderDeliveryArtifactDTO,
    TaskContractDraftDTO,
)
from ds_agent.domain.entities.messages import LLMResponse, Usage
from ds_agent.domain.entities.provider_models import ModelInfo
from ds_agent.infrastructure.task_contract_container import build_task_contract_container


@dataclass
class FakeProvider:
    response_text: str
    model_id: str = "gpt-5.4"
    provider_name: str = "openai"
    last_messages: list | None = None
    last_kwargs: dict[str, object] | None = None

    async def chat(self, messages, **kwargs):
        self.last_messages = messages
        self.last_kwargs = kwargs
        return LLMResponse(content=self.response_text, model=self.model_id, usage=Usage())

    async def count_tokens(self, messages) -> int:
        return 0

    def get_model_info(self) -> ModelInfo:
        return ModelInfo(
            model_id=self.model_id,
            provider=self.provider_name,
            display_name="Container Test Model",
            max_context_tokens=128_000,
            max_output_tokens=4096,
        )


def test_task_contract_container_renders_with_injected_llm_provider(tmp_path: Path) -> None:
    workspace_dir = str(tmp_path / "workspace")
    provider = FakeProvider(
        """
        {
          "blocks": [
            {
              "section": "summary",
              "title": "Summary",
              "body_md": "Premium cohorts drove the churn spike.",
              "citations": ["lineage-1"]
            }
          ],
          "overall_tone": "actionable",
          "flagged_claims": []
        }
        """
    )
    container = build_task_contract_container(workspace_dir, llm_provider=provider)
    created = container.create.execute(
        TaskContractDraftDTO(
            session_id="session-1",
            contract_type="churn_analysis",
            business_goal="Reduce churn",
            goal_brief={
                "business_question": "Why is churn increasing?",
                "ds_problem_statement": "Prioritize interventions",
                "comparison_baseline": "Current playbook",
                "decision_to_make": "Approve retention sprint",
                "expected_effort": "M",
            },
            required_deliverables=[
                {"type": "pm_action_memo", "audience": "pm", "format": "markdown"}
            ],
        )
    )
    built = container.build_delivery_pack.execute(
        BuildDeliveryPackDTO(task_id=created["task_id"], audiences=["pm"])
    )

    rendered = container.render_delivery_artifact.execute(
        RenderDeliveryArtifactDTO(
            task_id=created["task_id"],
            artifact_id=built["artifact_ids"][0],
            analysis={"summary": "Premium cohorts drove the churn spike."},
            output_dir=str(tmp_path / "artifacts"),
        )
    )

    assert Path(rendered["output_path"]).exists()
    assert rendered["verifier_status"] == "pass"
    assert provider.last_messages is not None
    assert provider.last_kwargs is not None
    assert provider.last_kwargs["response_format"] == {"type": "json_object"}


def test_task_contract_container_renders_audit_pdf_artifact(tmp_path: Path) -> None:
    workspace_dir = str(tmp_path / "workspace")
    container = build_task_contract_container(workspace_dir)
    created = container.create.execute(
        TaskContractDraftDTO(
            session_id="session-audit",
            contract_type="audit_review",
            business_goal="Provide audit evidence",
            goal_brief={
                "business_question": "Can the model be audited end-to-end?",
                "ds_problem_statement": "Assemble the audit trail",
                "comparison_baseline": "Current controls",
                "decision_to_make": "Approve production retention",
                "expected_effort": "S",
            },
            required_deliverables=[
                {"type": "audit_trail", "audience": "auditor", "format": "pdf"}
            ],
        )
    )
    built = container.build_delivery_pack.execute(
        BuildDeliveryPackDTO(
            task_id=created["task_id"],
            audiences=["auditor"],
            signature="signed-blob",
            tenant="deloitte",
        )
    )

    rendered = container.render_delivery_artifact.execute(
        RenderDeliveryArtifactDTO(
            task_id=created["task_id"],
            artifact_id=built["artifact_ids"][0],
            analysis={
                "summary": "Audit trail prepared",
                "data_provenance": ["warehouse.curated_churn", "feature_store.v2"],
                "access_log": ["etl-run-1", "approval-2"],
                "policy_compliance": ["retention label verified"],
                "approval_chain": ["owner approved"],
                "lineage": ["lineage-1"],
            },
            output_dir=str(tmp_path / "audit-artifacts"),
        )
    )

    assert rendered["format"] == "pdf"
    assert Path(rendered["output_path"]).exists()
    extracted = PdfReader(rendered["output_path"]).pages[0].extract_text()
    assert "Read-only compliance artifact." in extracted
    assert "Audit Trail" in extracted


def test_task_contract_container_applies_theme_to_exec_brief(tmp_path: Path) -> None:
    workspace_dir = str(tmp_path / "workspace-theme")
    container = build_task_contract_container(workspace_dir)
    created = container.create.execute(
        TaskContractDraftDTO(
            session_id="session-theme",
            contract_type="exec_briefing",
            business_goal="Brief leadership",
            goal_brief={
                "business_question": "What should leadership decide?",
                "ds_problem_statement": "Summarize the churn risk",
                "comparison_baseline": "Current retention baseline",
                "decision_to_make": "Approve intervention budget",
                "expected_effort": "S",
            },
            required_deliverables=[
                {"type": "exec_brief", "audience": "executive", "format": "pptx"}
            ],
        )
    )
    built = container.build_delivery_pack.execute(
        BuildDeliveryPackDTO(
            task_id=created["task_id"],
            audiences=["executive"],
            global_context={"theme_id": "deloitte_v1"},
        )
    )

    rendered = container.render_delivery_artifact.execute(
        RenderDeliveryArtifactDTO(
            task_id=created["task_id"],
            artifact_id=built["artifact_ids"][0],
            analysis={
                "summary": "Premium cohorts drove the churn spike.",
                "situation": "Churn increased in premium cohorts.",
                "finding": "Price sensitivity rose after the April change.",
                "impact": "Revenue risk is concentrated in two segments.",
                "recommendation": "Reverse the pricing step for at-risk cohorts.",
                "decision_needed": "Approve the targeted rollback.",
            },
            output_dir=str(tmp_path / "theme-artifacts"),
        )
    )

    title_run = (
        Presentation(rendered["output_path"]).slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    )
    assert title_run.font.name == "Open Sans"
    assert str(title_run.font.color.rgb) == "86BC25"

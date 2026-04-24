from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation

from ds_agent.domain.entities.delivery_pack import (
    ArtifactFormat,
    ArtifactType,
    AudienceKind,
    ContentPolicy,
    DeliveryArtifact,
    DeliveryPack,
    DeliveryTemplate,
    NarrativeBlocks,
    NarrativeVerification,
)
from ds_agent.infrastructure.exporters.pptx_exporter import PptxExporter


def test_pptx_exporter_creates_slides_and_footer(tmp_path: Path) -> None:
    template_path = tmp_path / "master.pptx"
    Presentation().save(template_path)
    artifact = DeliveryArtifact(
        artifact_id="art-exec",
        type=ArtifactType.EXEC_BRIEF,
        audience=AudienceKind.EXECUTIVE,
        format=ArtifactFormat.PPTX,
        content_policy=ContentPolicy(
            structure=["situation", "impact"],
            max_pages=3,
            chart_count_range=(1, 2),
            tone="decisive",
            technical_detail="minimal",
        ),
        template_ref="tpl/exec_brief/v3",
    )
    pack = DeliveryPack(
        pack_id="DP-1",
        task_id="TC-2026-001",
        source_analysis_id="fa-1",
        confidence=0.82,
        generated_at=datetime(2026, 4, 16, tzinfo=UTC),
        artifacts=[artifact],
    )
    narrative = NarrativeBlocks.model_validate(
        {
            "blocks": [
                {"section": "situation", "title": "Situation", "body_md": "- Churn is rising"},
                {"section": "impact", "title": "Impact", "body_md": "- Revenue at risk"},
            ],
            "overall_tone": "decisive",
            "flagged_claims": ["manual review"],
        }
    )

    output = PptxExporter().export(
        artifact=artifact,
        analysis={"summary": "Revenue at risk"},
        narrative=narrative,
        charts=[],
        verification=NarrativeVerification(
            report_id="vr-1",
            flagged_claims=["manual review"],
        ),
        template=DeliveryTemplate(
            template_ref=artifact.template_ref,
            family="exec_brief",
            format=ArtifactFormat.PPTX,
            title="Executive Brief",
            default_structure=("situation", "impact"),
            template_path=str(template_path),
        ),
        output_dir=tmp_path,
        pack=pack,
    )

    generated = Presentation(str(output))
    assert len(generated.slides) == 2
    footer_texts = []
    for shape in generated.slides[0].shapes:
        text = getattr(shape, "text", "")
        if "confidence=0.82" in text:
            footer_texts.append(text)
    assert footer_texts

    preview_manifest = json.loads(output.with_suffix(".preview.json").read_text(encoding="utf-8"))
    assert preview_manifest["artifact_id"] == "art-exec"
    assert preview_manifest["slide_count"] == 2
    assert preview_manifest["slides"][0]["title"] == "Situation"

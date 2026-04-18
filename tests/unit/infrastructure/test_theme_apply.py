from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation

from ds_agent.domain.entities.delivery_pack import (
    ArtifactFormat,
    ArtifactType,
    AudienceKind,
    ContentPolicy,
    DeliveryArtifact,
    DeliveryPack,
    DeliveryTemplate,
    NarrativeBlocks,
    NarrativeVerification,
)
from ds_agent.infrastructure.exporters.pptx_exporter import PptxExporter
from ds_agent.infrastructure.exporters.theme_loader import FileSystemThemeLoader


def _artifact() -> DeliveryArtifact:
    return DeliveryArtifact(
        artifact_id="art-exec",
        type=ArtifactType.EXEC_BRIEF,
        audience=AudienceKind.EXECUTIVE,
        format=ArtifactFormat.PPTX,
        content_policy=ContentPolicy(
            structure=["summary", "impact"],
            max_pages=3,
            chart_count_range=(1, 2),
            tone="decisive",
            technical_detail="minimal",
        ),
        template_ref="tpl/exec_brief/v3",
    )


def _narrative() -> NarrativeBlocks:
    return NarrativeBlocks.model_validate(
        {
            "blocks": [
                {"section": "summary", "title": "Summary", "body_md": "- Revenue at risk"}
            ],
            "overall_tone": "decisive",
            "flagged_claims": [],
        }
    )


def test_theme_loader_resolves_explicit_theme_and_tenant_default() -> None:
    loader = FileSystemThemeLoader()

    explicit = loader.resolve(
        pack=DeliveryPack(
            pack_id="DP-1",
            task_id="TC-1",
            generated_at=datetime(2026, 4, 16, tzinfo=UTC),
            global_context={"theme_id": "deloitte_v1"},
            artifacts=[],
        )
    )
    tenant_default = loader.resolve(
        pack=DeliveryPack(
            pack_id="DP-2",
            task_id="TC-2",
            generated_at=datetime(2026, 4, 16, tzinfo=UTC),
            tenant="deloitte",
            artifacts=[],
        )
    )

    assert explicit.theme_id == "deloitte_v1"
    assert explicit.primary_color == "#86BC25"
    assert tenant_default.theme_id == "deloitte_v1"


def test_pptx_exporter_applies_theme_from_pack_context(tmp_path: Path) -> None:
    master = tmp_path / "master.pptx"
    Presentation().save(master)
    artifact = _artifact()
    pack = DeliveryPack(
        pack_id="DP-3",
        task_id="TC-3",
        generated_at=datetime(2026, 4, 16, tzinfo=UTC),
        tenant="deloitte",
        global_context={"theme_id": "deloitte_v1"},
        artifacts=[artifact],
    )

    output = PptxExporter(theme_loader=FileSystemThemeLoader()).export(
        artifact=artifact,
        analysis={"summary": "Revenue at risk"},
        narrative=_narrative(),
        charts=[],
        verification=NarrativeVerification(report_id="vr-1"),
        template=DeliveryTemplate(
            template_ref=artifact.template_ref,
            family="exec_brief",
            format=ArtifactFormat.PPTX,
            title="Executive Brief",
            default_structure=("summary", "impact"),
            template_path=str(master),
        ),
        output_dir=tmp_path,
        pack=pack,
    )

    slide = Presentation(str(output)).slides[0]
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert title_run.font.name == "Open Sans"
    assert str(title_run.font.color.rgb) == "86BC25"

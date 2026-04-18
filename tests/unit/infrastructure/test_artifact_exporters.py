"""Tests for workspace-file exporters (P1-12)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from ds_agent.infrastructure.artifact.exporters import (
    ExportError,
    ExportFormat,
    export_file,
    supported_formats,
)


@pytest.fixture
def markdown_source(tmp_path: Path) -> Path:
    path = tmp_path / "report.md"
    path.write_text(
        "\n".join(
            [
                "# Title",
                "",
                "Intro paragraph with **bold** and _italic_.",
                "",
                "## Findings",
                "",
                "- first bullet",
                "- second bullet",
                "",
                "```python",
                "print('hello')",
                "```",
                "",
                "| col a | col b |",
                "|-------|-------|",
                "| 1 | 2 |",
                "| 3 | 4 |",
            ]
        ),
        encoding="utf-8",
    )
    return path


@pytest.fixture
def csv_source(tmp_path: Path) -> Path:
    path = tmp_path / "data.csv"
    path.write_text("a,b,c\n1,2.5,x\n3,4,y\n", encoding="utf-8")
    return path


@pytest.fixture
def notebook_source(tmp_path: Path) -> Path:
    path = tmp_path / "analysis.ipynb"
    path.write_text(
        json.dumps(
            {
                "nbformat": 4,
                "nbformat_minor": 5,
                "cells": [
                    {"cell_type": "markdown", "source": ["# Notebook"]},
                    {
                        "cell_type": "code",
                        "source": ["x = 1"],
                        "outputs": [
                            {
                                "name": "stdout",
                                "output_type": "stream",
                                "text": "done",
                            }
                        ],
                    },
                ],
            }
        ),
        encoding="utf-8",
    )
    return path


class TestSupportedFormats:
    def test_markdown_supports_pdf_docx_html(self) -> None:
        targets = supported_formats(".md")
        assert ExportFormat.PDF in targets
        assert ExportFormat.DOCX in targets
        assert ExportFormat.HTML in targets
        assert ExportFormat.XLSX not in targets

    def test_csv_supports_xlsx_html(self) -> None:
        targets = supported_formats(".csv")
        assert ExportFormat.XLSX in targets
        assert ExportFormat.HTML in targets
        assert ExportFormat.DOCX not in targets

    def test_unknown_extension_returns_empty(self) -> None:
        assert supported_formats(".exe") == []


class TestMarkdownExport:
    def test_html_renders_full_document(
        self, markdown_source: Path, tmp_path: Path
    ) -> None:
        output = tmp_path / "out.html"
        result = export_file(markdown_source, output, ExportFormat.HTML)
        content = result.output_path.read_text(encoding="utf-8")
        assert ">Title</h1>" in content
        assert "<table>" in content
        assert "<!DOCTYPE html>" in content

    def test_pdf_writes_intermediate_html(
        self, markdown_source: Path, tmp_path: Path
    ) -> None:
        output = tmp_path / "out.pdf.html"
        result = export_file(markdown_source, output, ExportFormat.PDF)
        assert result.intermediate_html_path == output
        assert ">Title</h1>" in output.read_text(encoding="utf-8")

    def test_docx_builds_valid_file(
        self, markdown_source: Path, tmp_path: Path
    ) -> None:
        from docx import Document

        output = tmp_path / "out.docx"
        export_file(markdown_source, output, ExportFormat.DOCX)
        doc = Document(str(output))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "Title" in headings
        assert len(doc.tables) == 1

    def test_cannot_export_markdown_to_xlsx(
        self, markdown_source: Path, tmp_path: Path
    ) -> None:
        with pytest.raises(ExportError):
            export_file(markdown_source, tmp_path / "x.xlsx", ExportFormat.XLSX)


class TestTabularExport:
    def test_xlsx_preserves_numeric_types(
        self, csv_source: Path, tmp_path: Path
    ) -> None:
        from openpyxl import load_workbook

        output = tmp_path / "out.xlsx"
        export_file(csv_source, output, ExportFormat.XLSX)
        wb = load_workbook(output)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == ("a", "b", "c")
        assert rows[1] == (1, 2.5, "x")

    def test_html_table_includes_header_cells(
        self, csv_source: Path, tmp_path: Path
    ) -> None:
        output = tmp_path / "out.html"
        export_file(csv_source, output, ExportFormat.HTML)
        content = output.read_text(encoding="utf-8")
        assert "<th>a</th>" in content
        assert "<td>1</td>" in content


class TestNotebookExport:
    def test_ipynb_passthrough_copies_content(
        self, notebook_source: Path, tmp_path: Path
    ) -> None:
        output = tmp_path / "copy.ipynb"
        export_file(notebook_source, output, ExportFormat.IPYNB)
        assert json.loads(output.read_text(encoding="utf-8"))["nbformat"] == 4

    def test_notebook_html_includes_code_and_outputs(
        self, notebook_source: Path, tmp_path: Path
    ) -> None:
        output = tmp_path / "nb.html"
        export_file(notebook_source, output, ExportFormat.HTML)
        content = output.read_text(encoding="utf-8")
        assert ">Notebook</h1>" in content
        assert 'class="code-cell"' in content
        assert "done" in content


class TestKoreanContent:
    """Korean / multi-byte content survives the export pipeline (P1-13 alignment)."""

    def test_markdown_html_preserves_korean_heading_and_body(self, tmp_path: Path) -> None:
        source = tmp_path / "보고서.md"
        source.write_text(
            "# 분석 결과\n\n주요 발견 사항은 다음과 같다.\n",
            encoding="utf-8",
        )
        output = tmp_path / "out.html"
        export_file(source, output, ExportFormat.HTML)
        content = output.read_text(encoding="utf-8")
        assert ">분석 결과</h1>" in content
        assert "주요 발견 사항은 다음과 같다" in content

    def test_csv_xlsx_preserves_korean_cells(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        source = tmp_path / "고객.csv"
        source.write_text("이름,점수\n홍길동,95\n김영희,88\n", encoding="utf-8")
        output = tmp_path / "out.xlsx"
        export_file(source, output, ExportFormat.XLSX)
        wb = load_workbook(output)
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[0] == ("이름", "점수")
        assert rows[1] == ("홍길동", 95)


class TestErrorPaths:
    def test_missing_source_raises(self, tmp_path: Path) -> None:
        with pytest.raises(ExportError):
            export_file(
                tmp_path / "nope.md", tmp_path / "out.html", ExportFormat.HTML
            )

    def test_directory_source_rejected(self, tmp_path: Path) -> None:
        subdir = tmp_path / "sub"
        subdir.mkdir()
        with pytest.raises(ExportError):
            export_file(subdir, tmp_path / "out.html", ExportFormat.HTML)

    def test_unsupported_combination_rejected(
        self, csv_source: Path, tmp_path: Path
    ) -> None:
        with pytest.raises(ExportError):
            export_file(csv_source, tmp_path / "x.docx", ExportFormat.DOCX)
